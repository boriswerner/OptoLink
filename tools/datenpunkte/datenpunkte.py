"""Datenpunkte der Heizung ermitteln: Kandidaten sammeln, an der Heizung roh lesen, auswerten.

Schritte (Ausgaben in docs/datenpunkte/lokal/, per .gitignore ausgeschlossen, da echte Messwerte):
  python datenpunkte.py kandidaten <addressen.xlsx>   -> kandidaten.json
  python datenpunkte.py scan [--host optolink] [--neu] -> scan.json (liest über /api/raw, 16 Adressen je
                                                         Paket; ein abgebrochener Scan wird fortgesetzt)
  python datenpunkte.py auswerten                     -> datenpunkte-vscotho1.csv, uebersicht.md

Abhängigkeit: openpyxl (nur für "kandidaten"). Der Scan liest nur, er schreibt nie in die Heizung.
"""
import argparse
import csv
import datetime
import json
import os
import re
import struct
import sys
import time
import urllib.error
import urllib.request

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "docs", "datenpunkte", "lokal")   # nicht einchecken: enthält echte Werte
MODEL = "VScotHO1"
MAX_LEN = 9  # MAX_DP_LENGTH von VitoWiFi

# VitoWiFi-Typen: Länge in Bytes
TYPE_LEN = {"DPTemp": 2, "DPTempS": 1, "DPMode": 1, "DPStat": 1, "DPCount": 4, "DPCountS": 2,
            "DPHours": 4, "DPCoP": 1, "DPPercent": 1, "DPRaw": 9}


# ---------------------------------------------------------------- Kandidaten

def type_from_factor(factor, length):
    """Leitet aus "Values Factor" und Bytes der Excel-Tabelle einen VitoWiFi-Typ ab."""
    f = factor.replace(" ", "").lower()
    if f in ("/10",):
        return "DPCoP" if length == 1 else "DPTemp"
    if f == "/2":
        return "DPPercent"
    if f.startswith("0..1") and length == 1:
        return "DPStat"
    if f == "sek":
        return "DPHours"
    if length == 4:
        return "DPCount"
    if length == 2:
        return "DPCountS"
    if length == 1:
        return "DPMode"
    return "DPRaw"


def split_cell(value):
    return [p for p in re.split(r"[\s/]+(?=0x)|\s+(?=0x)|\n", str(value).strip()) if p] if value else []


def candidates_from_excel(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = list(wb["Tabelle1"].iter_rows(values_only=True))
    models = [(i, v) for i, v in enumerate(rows[0]) if v]
    fields = rows[1]
    out = []
    section = ""
    for r in rows[2:]:
        name = (r[0] or "").strip() if isinstance(r[0], str) else (str(r[0]) if r[0] else "")
        if name and not any(r[2:]):
            section = name
            continue
        for idx, (start, model) in enumerate(models):
            end = models[idx + 1][0] if idx + 1 < len(models) else len(fields)
            block = {fields[c]: r[c] for c in range(start, min(end, len(r))) if fields[c]}
            addrs = re.findall(r"0x[0-9A-Fa-f]{4}", str(block.get("Addr") or ""))
            if not addrs:
                continue
            factors = str(block.get("Values Factor") or "").split() if len(addrs) > 1 else [str(block.get("Values Factor") or "")]
            sizes = str(block.get("Bytes") or "").split()
            for k, a in enumerate(addrs):
                factor = factors[k] if k < len(factors) else (factors[0] if factors else "")
                size = sizes[k] if k < len(sizes) else (sizes[0] if sizes else "")
                length = int(size) if size.isdigit() else 0
                if name == "Gerätekennung":
                    factor, length = "", 2
                out.append({"addr": a.upper().replace("0X", "0x"), "name": name, "short": str(r[1] or ""),
                            "section": section, "source": "Tabelle1 " + model, "model": model,
                            "factor": factor, "len": length,
                            "type": str(block.get("Datapoint Type") or "") or type_from_factor(factor, length or 2)})
    # eigene Tests (Blatt "Addressen")
    for r in list(wb["Addressen"].iter_rows(values_only=True))[1:]:
        addr = str(r[5] or "")
        if not re.fullmatch(r"0x[0-9A-Fa-f]{4}", addr):
            continue
        typ = str(r[1] or "")
        out.append({"addr": addr.upper().replace("0X", "0x"), "name": str(r[3] or r[2] or ""), "short": "",
                    "section": "", "source": "Addressen", "model": "", "factor": "",
                    "len": TYPE_LEN.get(typ, 0), "type": typ,
                    "user_ok": str(r[7] or ""), "user_value": str(r[8] or ""), "user_comment": str(r[9] or "")})
    return out


def candidates_from_firmware():
    """Aktive Datenpunkte aus src/viessmann.cpp (Name = MQTT-Topic)."""
    text = open(os.path.join(ROOT, "src", "viessmann.cpp"), encoding="utf-8").read()
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.S)
    out = []
    for m in re.finditer(r'^\s*(DP\w+)\s+\w+\("([^"]+)",\s*"[^"]*",\s*(0x[0-9A-Fa-f]+)\)', text, re.M):
        typ, name, addr = m.groups()
        length = 9 if typ == "DPRaw" else TYPE_LEN.get(typ, 2)
        out.append({"addr": "0x%04X" % int(addr, 16), "name": name, "short": "", "section": "",
                    "source": "Firmware", "model": "", "factor": "", "len": length, "type": typ, "topic": name})
    return out


def cmd_kandidaten(args):
    cands = candidates_from_excel(args.xlsx) + candidates_from_firmware()
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "kandidaten.json"), "w", encoding="utf-8") as f:
        json.dump(cands, f, ensure_ascii=False, indent=1)
    addrs = {c["addr"] for c in cands}
    print(f"{len(cands)} Einträge, {len(addrs)} verschiedene Adressen")


# ---------------------------------------------------------------- Scan

def http(host, path, method="GET", timeout=10):
    req = urllib.request.Request(f"http://{host}{path}", method=method, data=b"" if method == "POST" else None)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8", "replace")


BATCH = 16  # SCAN_BATCH der Firmware


def read_batch(host, parts):
    """Liest bis zu BATCH (Adresse, Länge)-Paare in einem Rutsch. Liefert {Adresse: hex|None}."""
    query = ",".join(f"0x{a:04X}:{n}" for a, n in parts)
    for _ in range(100):  # "belegt": reguläre Abfrage läuft evtl. noch
        status, body = http(host, f"/api/raw?list={query}")
        if status == 202:
            rid = json.loads(body)["id"]
            break
        if status == 409 and "belegt" in body:
            time.sleep(0.2)
            continue
        raise RuntimeError(f"HTTP {status} {body}")
    else:
        raise RuntimeError("dauerhaft belegt")
    deadline = time.time() + 5 * len(parts) + 10  # 5 s Timeout je Adresse, falls nichts antwortet
    while time.time() < deadline:
        res = json.loads(http(host, "/api/raw/result")[1])
        if res["id"] == rid and not res["busy"]:
            return {int(r["addr"], 16): (r["hex"] if r["state"] == "done" else None) for r in res["results"]}
        time.sleep(0.1)
    raise RuntimeError("keine Rückmeldung")


def cmd_scan(args):
    cands = json.load(open(os.path.join(OUT, "kandidaten.json"), encoding="utf-8"))
    lengths = {}
    for c in cands:
        lengths[c["addr"]] = max(lengths.get(c["addr"], 0), c["len"] or 4)
    path = os.path.join(OUT, "scan.json")
    results = {}
    if os.path.exists(path) and not args.neu:  # Zwischenstand fortsetzen
        results = json.load(open(path, encoding="utf-8"))["werte"]
    todo = sorted((a for a in lengths if a not in results), key=lambda a: int(a, 16))
    # lange Blöcke in 9-Byte-Stücke zerlegen
    parts = [(int(a, 16) + off, min(MAX_LEN, lengths[a] - off), a) for a in todo for off in range(0, lengths[a], MAX_LEN)]
    print(f"{len(todo)} Adressen offen ({len(parts)} Lesevorgänge), {len(results)} schon gelesen", flush=True)

    def save():
        with open(path, "w", encoding="utf-8") as f:
            json.dump({"zeit": datetime.datetime.now().isoformat(timespec="seconds"), "werte": results}, f, indent=1)

    status, body = http(args.host, "/api/scan/start", "POST")
    if status != 200:
        sys.exit(f"Scan-Start fehlgeschlagen: HTTP {status} {body}")
    pieces = {}
    try:
        for i in range(0, len(parts), BATCH):
            batch = parts[i:i + BATCH]
            got = read_batch(args.host, [(p[0], p[1]) for p in batch])
            for start, n, addr in batch:
                pieces.setdefault(addr, []).append(got.get(start))
            for addr in {p[2] for p in batch}:
                expected = len(range(0, lengths[addr], MAX_LEN))
                if len(pieces[addr]) == expected:
                    ok = all(h is not None for h in pieces[addr])
                    results[addr] = {"len": lengths[addr], "hex": "".join(pieces[addr]) if ok else None}
            save()
            failed = sum(1 for p in batch if got.get(p[0]) is None)
            print(f"{min(i + BATCH, len(parts)):3}/{len(parts)} Lesevorgänge, in diesem Paket ohne Antwort: {failed}", flush=True)
    finally:
        http(args.host, "/api/scan/stop", "POST")
        save()
    print(f"{len(results)} Adressen gelesen, {sum(1 for r in results.values() if r['hex'] is None)} ohne Antwort")


# ---------------------------------------------------------------- Auswertung

def decode(typ, raw):
    """Wert wie VitoWiFi ihn liefern würde (Little Endian)."""
    n = TYPE_LEN.get(typ, 0)
    if typ == "DPRaw" or len(raw) < n:
        return None
    b = raw[:n]
    if typ == "DPTemp":
        return struct.unpack("<h", b)[0] / 10
    if typ in ("DPTempS", "DPMode"):
        return b[0]
    if typ == "DPStat":
        return b[0]
    if typ == "DPCoP":
        return b[0] / 10
    if typ == "DPPercent":
        return b[0] / 2
    if typ == "DPCountS":
        return struct.unpack("<H", b)[0]
    if typ == "DPCount":
        return struct.unpack("<I", b)[0]
    if typ == "DPHours":
        return round(struct.unpack("<I", b)[0] / 3600, 1)
    return None


def invalid_marker(typ, raw):
    n = TYPE_LEN.get(typ, len(raw))
    b = raw[:n]
    if b and all(x == 0xFF for x in b):
        return "alle Bytes FF (Fehlwert/nicht belegt)"
    if typ == "DPTemp" and len(b) == 2 and struct.unpack("<h", b)[0] in (0x7FFF, -32768):
        return "7FFF/8000 (Fehlwert/Sensor fehlt)"
    return ""


def judge(typ, value, raw, names):
    """Bewertung: plausibel / fraglich / Fehlwert, mit Begründung."""
    if raw is None:
        return "keine Antwort", "Lesefehler/Timeout"
    marker = invalid_marker(typ, raw)
    if marker:
        return "Fehlwert", marker
    text = " ".join(names).lower()
    if typ == "DPRaw":
        if len(raw) >= 9 and raw[1] == 0x20 and 0 < (raw[2] >> 4) * 10 + (raw[2] & 15) < 100:
            return "plausibel", "Fehlereintrag mit BCD-Datum"
        return "fraglich", "Rohdaten"
    if value is None:
        return "fraglich", "nicht deutbar"
    if typ == "DPTemp":
        if "soll" in text:
            return ("plausibel", "") if 0 <= value <= 100 else ("fraglich", "Sollwert außerhalb 0..100")
        if value == 0:
            return "fraglich", "genau 0"
        return ("plausibel", "") if -40 <= value <= 150 else ("fraglich", "außerhalb -40..150 °C")
    if typ == "DPTempS":
        return ("plausibel", "") if 3 <= value <= 95 else ("fraglich", "Sollwert außerhalb 3..95")
    if typ == "DPStat":
        return ("plausibel", "") if value in (0, 1) else ("fraglich", "kein 0/1")
    if typ == "DPPercent":
        return ("plausibel", "") if 0 <= value <= 100 else ("fraglich", "außerhalb 0..100 %")
    if typ == "DPCoP":
        return ("plausibel", "") if 0.2 <= value <= 3.5 else ("fraglich", "Neigung außerhalb 0,2..3,5")
    if typ == "DPMode":
        if "0..1" in text:
            return ("plausibel", "") if value in (0, 1) else ("fraglich", "kein 0/1")
        return ("plausibel", "") if value <= 100 else ("fraglich", "Zustand > 100")
    if typ in ("DPCount", "DPCountS", "DPHours"):
        return ("plausibel", "") if value > 0 else ("fraglich", "genau 0")
    return "fraglich", ""


def variants(raw):
    if not raw:
        return ""
    v = [f"u8={raw[0]}", f"s8={struct.unpack('<b', raw[:1])[0]}"]
    if len(raw) >= 2:
        v += [f"u16={struct.unpack('<H', raw[:2])[0]}", f"s16/10={struct.unpack('<h', raw[:2])[0] / 10}",
              f"u16/100={struct.unpack('<H', raw[:2])[0] / 100}"]
    if len(raw) >= 4:
        v += [f"u32={struct.unpack('<I', raw[:4])[0]}"]
    return " ".join(v)


PRIORITY = {"Firmware": 0, "Addressen": 1, "Tabelle1 " + MODEL: 2}


def cmd_auswerten(args):
    cands = json.load(open(os.path.join(OUT, "kandidaten.json"), encoding="utf-8"))
    scan = json.load(open(os.path.join(OUT, "scan.json"), encoding="utf-8"))
    by_addr = {}
    for c in cands:
        by_addr.setdefault(c["addr"], []).append(c)
    rows = []
    for addr in sorted(by_addr, key=lambda a: int(a, 16)):
        cs = sorted(by_addr[addr], key=lambda c: PRIORITY.get(c["source"], 3))
        res = scan["werte"].get(addr, {})
        raw = bytes.fromhex(res["hex"]) if res.get("hex") else None
        names = [c["name"] for c in cs if c["name"]]
        # Typ-Kandidaten in Prioritätsreihenfolge; der erste plausible gewinnt
        types = []
        for c in cs:
            if c["type"] and c["type"] not in types:
                types.append(c["type"])
        best = None
        for t in types or ["DPRaw"]:
            value = decode(t, raw) if raw else None
            verdict, why = judge(t, value, raw, names)
            if best is None or (verdict == "plausibel" and best[2] != "plausibel"):
                best = (t, value, verdict, why)
        typ, value, verdict, why = best
        tab = next((c for c in cs if c["source"] == "Tabelle1 " + MODEL), None)
        name = (tab or next((c for c in cs if c["source"].startswith("Tabelle1")), None) or cs[0])["name"]
        topic = next((c.get("topic") for c in cs if c.get("topic")), "")
        user = next((c for c in cs if c["source"] == "Addressen"), None)
        hints = [why] if why else []
        if user and (user.get("user_ok") or user.get("user_comment")):
            hints.append(f"Test 2022: OK={user['user_ok'] or '-'} Wert={user['user_value'] or '-'} {user['user_comment']}".strip())
        rows.append({
            "Adresse": addr, "Länge": res.get("len", ""), "Klarname": name,
            "Weitere Bezeichnungen": "; ".join(sorted(set(names) - {name})),
            "Quellen": "; ".join(sorted({c["source"] for c in cs})),
            "MQTT-Topic (heute)": topic, "Rohwert (hex)": res.get("hex") or "",
            "Vorgeschlagener Typ": typ, "Wert": "" if value is None else value,
            "Bewertung": verdict, "Hinweis": " | ".join(hints), "Varianten": variants(raw),
            "Bestätigter Typ": "", "Kommentar": ""})
    path = os.path.join(OUT, "datenpunkte-vscotho1.csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(rows)
    counts = {}
    for r in rows:
        counts[r["Bewertung"]] = counts.get(r["Bewertung"], 0) + 1
    with open(os.path.join(OUT, "uebersicht.md"), "w", encoding="utf-8") as f:
        f.write(f"# Datenpunkte {MODEL}\n\nScan vom {scan['zeit']}, {len(rows)} Adressen: "
                + ", ".join(f"{v} {k}" for k, v in sorted(counts.items())) + ".\n\n"
                "Vollständige Liste mit Rohwerten und Spalten für eigene Angaben: "
                "[datenpunkte-vscotho1.csv](datenpunkte-vscotho1.csv).\n\n"
                "## Plausible Datenpunkte\n\n| Adresse | Klarname | Typ | Wert | MQTT-Topic (heute) |\n|---|---|---|---|---|\n")
        for r in rows:
            if r["Bewertung"] == "plausibel":
                f.write(f"| {r['Adresse']} | {r['Klarname']} | {r['Vorgeschlagener Typ']} | {r['Wert']} | {r['MQTT-Topic (heute)']} |\n")
    print(f"{len(rows)} Adressen ausgewertet: {counts} -> {path}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    k = sub.add_parser("kandidaten")
    k.add_argument("xlsx")
    s = sub.add_parser("scan")
    s.add_argument("--host", default="optolink")
    s.add_argument("--neu", action="store_true", help="Zwischenstand verwerfen und alles neu lesen")
    sub.add_parser("auswerten")
    args = p.parse_args()
    {"kandidaten": cmd_kandidaten, "scan": cmd_scan, "auswerten": cmd_auswerten}[args.cmd](args)


if __name__ == "__main__":
    main()
